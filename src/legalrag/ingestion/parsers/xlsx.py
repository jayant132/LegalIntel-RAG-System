"""XLSX parser: one typed reader, two consumers.

``read_sheets`` returns typed cell values (numbers stay numbers) so the SQL path for
aggregation questions can reuse it. ``parse_xlsx`` renders the same data as markdown tables
and note blocks for retrieval.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText

from legalrag.ingestion.manifest import DocumentRecord
from legalrag.ingestion.models import Block, BlockKind, ParsedDocument
from legalrag.ingestion.parsers._common import to_markdown_table

CellValue = str | int | float | bool | date | datetime | None


@dataclass(frozen=True)
class SheetData:
    """One worksheet: a header, data rows, and free-text note rows (e.g. footnotes)."""

    name: str
    header: tuple[str, ...]
    rows: tuple[tuple[CellValue, ...], ...]
    notes: tuple[str, ...]


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _coerce(value: object) -> CellValue:
    """Narrow openpyxl's wide cell type to ours; reject anything we don't understand."""
    if value is None or isinstance(value, str | bool | int | float | date | datetime):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, CellRichText):
        return str(value)  # plain text of a rich-text cell
    raise ValueError(f"unsupported cell value type: {type(value).__name__}")


def read_sheets(path: Path) -> list[SheetData]:
    """Read every sheet, failing loudly on anything that could silently lose data.

    Raises ``ValueError`` for hidden sheets, formulas with no cached result, blank header
    cells, values outside the header's columns, and unsupported cell types (e.g. times).
    """
    values_wb = load_workbook(path, data_only=True)
    formulas_wb = load_workbook(path, data_only=False)
    sheets: list[SheetData] = []
    for ws in values_wb.worksheets:
        if ws.sheet_state != "visible":
            raise ValueError(f"{path.name}: hidden sheet {ws.title!r}; review before ingesting")
        for row in formulas_wb[ws.title].iter_rows():
            for cell in row:
                if cell.data_type == "f" and ws[cell.coordinate].value is None:
                    raise ValueError(
                        f"{path.name}!{ws.title}!{cell.coordinate}: formula with no cached value"
                    )

        header: tuple[str, ...] | None = None
        rows: list[tuple[CellValue, ...]] = []
        notes: list[str] = []
        for raw in ws.iter_rows(values_only=True):
            if all(_is_blank(v) for v in raw):
                continue
            if header is None:
                head = list(raw)
                while head and _is_blank(head[-1]):
                    head.pop()
                if any(_is_blank(c) for c in head):
                    raise ValueError(f"{path.name}!{ws.title}: blank cell in header row")
                header = tuple(str(c).strip() for c in head)
                continue
            width = len(header)
            if any(not _is_blank(v) for v in raw[width:]):
                raise ValueError(f"{path.name}!{ws.title}: value outside the header's columns")
            cells = [_coerce(v) for v in raw[:width]]
            filled = [v for v in cells if not _is_blank(v)]
            if width > 1 and len(filled) == 1 and isinstance(cells[0], str):
                notes.append(cells[0].strip())  # a footnote, not a data row
                continue
            rows.append(tuple(cells))
        if header is not None:
            sheets.append(SheetData(ws.title, header, tuple(rows), tuple(notes)))
    return sheets


def _format(value: CellValue) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_xlsx(record: DocumentRecord, path: Path) -> ParsedDocument:
    """One markdown TABLE block per sheet, plus one TEXT block per footnote."""
    blocks: list[Block] = []
    for sheet in read_sheets(path):
        if sheet.rows:
            table = to_markdown_table(
                [sheet.header, *[[_format(v) for v in row] for row in sheet.rows]]
            )
            blocks.append(Block(kind=BlockKind.TABLE, text=f"Sheet: {sheet.name}\n\n{table}"))
        blocks.extend(
            Block(kind=BlockKind.TEXT, text=f"Sheet: {sheet.name}\nNote: {note}")
            for note in sheet.notes
        )
    return ParsedDocument(doc_id=record.doc_id, blocks=tuple(blocks))
