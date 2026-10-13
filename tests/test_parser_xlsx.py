from collections.abc import Callable
from datetime import time
from pathlib import Path

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from legalrag.domain import DocType
from legalrag.ingestion.manifest import DocumentRecord
from legalrag.ingestion.models import BlockKind
from legalrag.ingestion.parsers.xlsx import parse_xlsx, read_sheets

MakeRecord = Callable[..., DocumentRecord]


def _save(tmp_path: Path, wb: Workbook) -> Path:
    p = tmp_path / "f.xlsx"
    wb.save(p)
    return p


def _simple_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Reg"
    ws.append(["id", "value", "ratio"])
    ws.append(["A", 1450000.0, 99.97])
    ws.append(["B", None, 5])
    ws.append([])
    ws.append(["Note: figures are preliminary"])
    return wb


def test_table_and_footnote_become_separate_blocks(tmp_path: Path, make_record: MakeRecord) -> None:
    doc = parse_xlsx(make_record(DocType.REGISTER, "xlsx"), _save(tmp_path, _simple_workbook()))
    assert [b.kind for b in doc.blocks] == [BlockKind.TABLE, BlockKind.TEXT]
    assert "| A | 1450000 | 99.97 |" in doc.blocks[0].text  # whole floats lose the ".0"
    assert "| B |  | 5 |" in doc.blocks[0].text  # None -> empty cell
    assert doc.blocks[1].text == "Sheet: Reg\nNote: Note: figures are preliminary"


def test_read_sheets_keeps_numbers_as_numbers(tmp_path: Path) -> None:
    (sheet,) = read_sheets(_save(tmp_path, _simple_workbook()))
    assert sheet.header == ("id", "value", "ratio")
    assert sheet.rows[0] == ("A", 1450000.0, 99.97)
    assert sheet.notes == ("Note: figures are preliminary",)


def test_hidden_sheet_is_rejected(tmp_path: Path) -> None:
    wb = _simple_workbook()
    wb.create_sheet("secret").sheet_state = "hidden"
    with pytest.raises(ValueError, match="hidden sheet"):
        read_sheets(_save(tmp_path, wb))


def test_formula_without_cached_value_is_rejected(tmp_path: Path) -> None:
    wb = _simple_workbook()
    ws = wb["Reg"]
    ws["B3"] = "=1+1"  # openpyxl writes no cached result
    with pytest.raises(ValueError, match="formula with no cached value"):
        read_sheets(_save(tmp_path, wb))


def test_blank_header_cell_is_rejected(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["id", None, "x"])
    ws.append(["A", 1, 2])
    with pytest.raises(ValueError, match="blank cell in header"):
        read_sheets(_save(tmp_path, wb))


def test_unsupported_cell_type_is_rejected(tmp_path: Path) -> None:
    wb = _simple_workbook()
    wb["Reg"]["B3"] = time(9, 30)
    with pytest.raises(ValueError, match="unsupported cell value type"):
        read_sheets(_save(tmp_path, wb))


def test_workbook_with_no_content_fails_loudly(tmp_path: Path, make_record: MakeRecord) -> None:
    with pytest.raises(ValidationError, match="no content"):
        parse_xlsx(make_record(DocType.REGISTER, "xlsx"), _save(tmp_path, Workbook()))


def test_real_contract_register(manifest: dict[str, DocumentRecord], data_dir: Path) -> None:
    rec = manifest["D16"]
    (sheet,) = read_sheets(data_dir / rec.path)
    assert len(sheet.rows) == 11
    first = dict(zip(sheet.header, sheet.rows[0], strict=True))
    assert first["contract_id"] == "C-001"
    assert first["notice_period_days"] == 120  # the stale value; A2 restored 180
    doc = parse_xlsx(rec, data_dir / rec.path)
    assert len(doc.blocks) == 1
    assert "Nordhaven Cloud Solutions A/S" in doc.text


def test_real_sla_report_has_two_tables_and_the_footnote(
    manifest: dict[str, DocumentRecord], data_dir: Path
) -> None:
    rec = manifest["D15"]
    doc = parse_xlsx(rec, data_dir / rec.path)
    kinds = [b.kind for b in doc.blocks]
    assert kinds == [BlockKind.TABLE, BlockKind.TEXT, BlockKind.TABLE]
    assert "| Nov 2024 | 99.12 | 99.95 | 10% |" in doc.blocks[0].text
    assert "Amendment 2" in doc.blocks[1].text
    assert "INC-53744" in doc.blocks[2].text
