from collections.abc import Callable
from pathlib import Path

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from pydantic import ValidationError

from legalrag.domain import DocType
from legalrag.ingestion.manifest import DocumentRecord
from legalrag.ingestion.models import BlockKind, ParsedDocument
from legalrag.ingestion.parsers.docx import parse_docx

MakeRecord = Callable[..., DocumentRecord]


def _parse(tmp_path: Path, make_record: MakeRecord, doc: DocxDocument) -> ParsedDocument:
    p = tmp_path / "f.docx"
    doc.save(str(p))
    return parse_docx(make_record(DocType.AMENDMENT, "docx"), p)


def _doc(*paragraphs: str) -> DocxDocument:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    return doc


def test_sections_group_heading_with_clauses(tmp_path: Path, make_record: MakeRecord) -> None:
    doc = _doc("TITLE", "1. Purpose", "Some text.", "2. Fees", "2.1 First.", "2.2 Second.")
    parsed = _parse(tmp_path, make_record, doc)
    assert [b.text for b in parsed.blocks] == [
        "TITLE",
        "1. Purpose\nSome text.",
        "2. Fees\n2.1 First.\n2.2 Second.",
    ]


def test_sentences_starting_with_a_number_are_not_headings(
    tmp_path: Path, make_record: MakeRecord
) -> None:
    doc = _doc("1. Purpose", "2. This is a full sentence that ends with a period.")
    parsed = _parse(tmp_path, make_record, doc)
    assert len(parsed.blocks) == 1


def test_page_break_ends_the_section_and_resets_the_heading(
    tmp_path: Path, make_record: MakeRecord
) -> None:
    doc = _doc("9. Governing law", "English law applies.")
    doc.add_page_break()
    doc.add_paragraph("LETTER OF VARIATION")
    doc.add_paragraph("Salary increases.")
    parsed = _parse(tmp_path, make_record, doc)
    assert [b.text for b in parsed.blocks] == [
        "9. Governing law\nEnglish law applies.",
        "LETTER OF VARIATION\nSalary increases.",
    ]


def test_heading_without_content_is_kept(tmp_path: Path, make_record: MakeRecord) -> None:
    parsed = _parse(tmp_path, make_record, _doc("1. Lonely heading", "2. Next", "Body."))
    assert [b.text for b in parsed.blocks] == ["1. Lonely heading", "2. Next\nBody."]


def test_table_stays_in_its_section_with_heading_and_caption(
    tmp_path: Path, make_record: MakeRecord
) -> None:
    doc = _doc("7. Records", "Keep records.", "Version history")
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip(table._cells, ["Version", "Date", "4.0", "2024"], strict=True):
        cell.text = text
    parsed = _parse(tmp_path, make_record, doc)
    assert parsed.blocks[0].text == "7. Records\nKeep records."
    assert parsed.blocks[1].kind is BlockKind.TABLE
    assert parsed.blocks[1].text.startswith("7. Records\nVersion history\n\n| Version | Date |")
    assert "| 4.0 | 2024 |" in parsed.blocks[1].text


def test_merged_cells_are_not_repeated(tmp_path: Path, make_record: MakeRecord) -> None:
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ["A", "B", "C"], strict=True):
        cell.text = text
    table.cell(1, 0).merge(table.cell(1, 1)).text = "merged"
    table.cell(1, 2).text = "z"
    parsed = _parse(tmp_path, make_record, doc)
    assert parsed.blocks[0].text.splitlines()[2] == "| merged | z |  |"


@pytest.mark.parametrize(
    ("tag", "message"),
    [("w:ins", "tracked insertion"), ("w:footnoteReference", "footnote")],
)
def test_constructs_that_would_lose_text_are_rejected(
    tmp_path: Path, make_record: MakeRecord, tag: str, message: str
) -> None:
    doc = _doc("1. Clause", "Body.")
    doc.paragraphs[1]._p.append(OxmlElement(tag))
    with pytest.raises(ValueError, match=message):
        _parse(tmp_path, make_record, doc)


def test_nested_table_is_rejected(tmp_path: Path, make_record: MakeRecord) -> None:
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    outer.cell(0, 0).add_table(rows=1, cols=1)
    with pytest.raises(ValueError, match="nested table"):
        _parse(tmp_path, make_record, doc)


def test_empty_document_fails_loudly(tmp_path: Path, make_record: MakeRecord) -> None:
    with pytest.raises(ValidationError, match="no content"):
        _parse(tmp_path, make_record, Document())


def test_real_amendment_2_sections(manifest: dict[str, DocumentRecord], data_dir: Path) -> None:
    rec = manifest["D05"]
    parsed = parse_docx(rec, data_dir / rec.path)
    assert len(parsed.blocks) == 7
    assert parsed.blocks[2].text.startswith("2. Termination for convenience\n2.1 Clause 12.1")


def test_real_employment_agreement_variation_is_a_separate_block(
    manifest: dict[str, DocumentRecord], data_dir: Path
) -> None:
    rec = manifest["D06"]
    parsed = parse_docx(rec, data_dir / rec.path)
    letter = parsed.blocks[-1]
    assert letter.text.startswith("LETTER OF VARIATION")
    assert "GBP 198,000" in letter.text
    assert "Governing law" not in letter.text  # not glued to clause 9
    tables = [b for b in parsed.blocks if b.kind is BlockKind.TABLE]
    assert len(tables) == 1
    assert tables[0].text.startswith("2. Remuneration\n\n| Scorecard element |")


def test_real_policy_v4_tables(manifest: dict[str, DocumentRecord], data_dir: Path) -> None:
    rec = manifest["D08"]
    parsed = parse_docx(rec, data_dir / rec.path)
    tables = [b for b in parsed.blocks if b.kind is BlockKind.TABLE]
    assert len(tables) == 2
    assert "Version history" in tables[1].text
    assert "GBP 75" in parsed.text


def test_real_board_minutes(manifest: dict[str, DocumentRecord], data_dir: Path) -> None:
    rec = manifest["D07"]
    parsed = parse_docx(rec, data_dir / rec.path)
    assert len(parsed.blocks) == 6
    assert "RESOLVED" in parsed.text
